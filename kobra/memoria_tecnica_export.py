# © 2026 Martín Viera. Todos los derechos reservados.

"""
MV Kobra AI · La memoria técnica, en los tres formatos que se piden
====================================================================
HTML para mandar por mail o publicar, Word para que alguien lo edite y le
agregue lo suyo, PDF para imprimir y llevar a una reunión.

Los tres salen del MISMO catálogo (`kobra/memoria_tecnica.py`). No hay una
plantilla por formato con el texto adentro: eso garantiza que a los tres meses
el PDF diga una cosa y el HTML otra, y que nadie se entere hasta que un cliente
compare los dos.

Por qué los renderizadores están acá y no en el catálogo
--------------------------------------------------------
El catálogo lo leen la API, la pantalla y los tests; los renderizadores solo
los usa el export. Separarlos evita que `kobra/memoria_tecnica.py` arrastre
`fpdf` y `python-docx` cada vez que la pantalla pide la lista — que es el 99%
de las veces.
"""
from __future__ import annotations

import html
import io

from kobra import memoria_tecnica as mt

TITULO = "MV Kobra AI · Memoria técnica del pipeline"
BAJADA = ("Qué hace el programa, en orden, contado dos veces: una para quien "
          "lee el código y otra para quien firma la compra.")

# Los dos registros, con el nombre que van a leer las dos audiencias.
_ETIQUETA = {
    "tecnico": "Técnico",
    "criollo": "En criollo",
    "por_que": "Por qué se hizo",
    "repercusion": "Cómo repercute aguas abajo",
    "limites": "Qué NO hace",
}


def _pie(datos_demo: bool) -> str:
    """La misma aclaración en los tres formatos.

    No es un disclaimer de trámite: la demo corre sobre datos sintéticos y las
    cifras de impacto son ilustrativas. Un documento técnico que no lo diga se
    lee como si los números fueran de un cliente real.
    """
    base = ("Documento generado por el propio programa a partir de "
            "kobra/memoria_tecnica.py: describe el pipeline tal como está "
            "implementado.")
    if datos_demo:
        base += (" Los datos de esta instalación son SINTÉTICOS y las cifras "
                 "de impacto son ILUSTRATIVAS, no resultados medidos en "
                 "producción.")
    return base


# --- HTML -------------------------------------------------------------------
def como_html(datos_demo: bool = True) -> str:
    """Un solo archivo, sin CSS externo ni JavaScript.

    Autocontenido a propósito: este HTML se manda adjunto por mail y se abre
    con doble clic desde el Escritorio. Una hoja de estilos por CDN lo deja sin
    formato en la mitad de esos casos, y un `<script>` hace que algunos
    clientes de correo lo bloqueen entero.
    """
    e_ = html.escape
    partes = [
        "<!doctype html>", '<html lang="es"><head><meta charset="utf-8">',
        f"<title>{e_(TITULO)}</title>",
        "<style>",
        "body{font:15px/1.6 system-ui,-apple-system,Segoe UI,Roboto,sans-serif;",
        "     max-width:52rem;margin:2rem auto;padding:0 1.2rem;color:#16202b}",
        "h1{font-size:1.7rem;margin:0 0 .3rem}",
        ".bajada{color:#5b6b7c;margin:0 0 2rem}",
        "article{border-top:1px solid #dfe6ee;padding:1.4rem 0}",
        "h2{font-size:1.15rem;margin:0 0 .1rem}",
        ".orden{color:#8a99a8;font-variant-numeric:tabular-nums;font-size:.85rem}",
        "dt{font-weight:600;margin-top:.9rem;font-size:.82rem;",
        "   text-transform:uppercase;letter-spacing:.04em;color:#41566b}",
        "dd{margin:.25rem 0 0}",
        ".criollo{background:#f4f8fc;border-left:3px solid #2f6fb0;",
        "         padding:.7rem .9rem;border-radius:0 4px 4px 0}",
        ".limites{color:#7a4a1e}",
        ".meta{margin-top:1rem;font-size:.82rem;color:#6b7b8b}",
        "code{background:#eef3f8;padding:.1rem .3rem;border-radius:3px;",
        "     font-size:.85em}",
        "footer{border-top:1px solid #dfe6ee;margin-top:2rem;padding-top:1rem;",
        "       font-size:.82rem;color:#6b7b8b}",
        "@media print{body{max-width:none;margin:0}article{break-inside:avoid}}",
        "</style></head><body>",
        f"<h1>{e_(TITULO)}</h1>",
        f'<p class="bajada">{e_(BAJADA)}</p>',
    ]
    for et in mt.etapas():
        partes += [
            "<article>",
            f'<div class="orden">Etapa {et.orden} de {len(mt.etapas())}</div>',
            f"<h2>{e_(et.titulo)}</h2>",
            "<dl>",
            f"<dt>{_ETIQUETA['tecnico']}</dt><dd>{e_(et.tecnico)}</dd>",
            f"<dt>{_ETIQUETA['criollo']}</dt>"
            f'<dd class="criollo">{e_(et.criollo)}</dd>',
            f"<dt>{_ETIQUETA['por_que']}</dt><dd>{e_(et.por_que)}</dd>",
            f"<dt>{_ETIQUETA['repercusion']}</dt><dd>{e_(et.repercusion)}</dd>",
        ]
        if et.limites:
            partes.append(f"<dt>{_ETIQUETA['limites']}</dt>"
                          f'<dd class="limites">{e_(et.limites)}</dd>')
        partes += [
            "</dl>",
            '<p class="meta">',
            f"<b>Entra:</b> {e_(et.entradas)}<br>",
            f"<b>Sale:</b> {e_(et.salidas)}<br>",
            "<b>Dónde vive:</b> "
            + " · ".join(f"<code>{e_(m)}</code>" for m in et.modulos),
            "</p></article>",
        ]
    partes += [f"<footer>{e_(_pie(datos_demo))}</footer>", "</body></html>"]
    return "\n".join(partes)


# --- Word -------------------------------------------------------------------
def como_docx(datos_demo: bool = True) -> bytes:
    """`.docx` de verdad, no un HTML con extensión cambiada.

    El atajo conocido es servir HTML con `Content-Type: application/msword`:
    Word lo abre, pero el archivo miente sobre lo que es, se rompe al editarlo
    en Google Docs o LibreOffice, y deja al que lo recibe explicando por qué su
    editor no lo toma. `python-docx` cuesta 250 KB en el instalador y devuelve
    un documento que cualquier editor entiende.
    """
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(TITULO, level=0)
    doc.add_paragraph(BAJADA)

    for et in mt.etapas():
        doc.add_heading(f"{et.orden}. {et.titulo}", level=1)
        for clave in ("tecnico", "criollo", "por_que", "repercusion", "limites"):
            texto = getattr(et, clave)
            if not texto:
                continue
            p = doc.add_paragraph()
            r = p.add_run(f"{_ETIQUETA[clave]}: ")
            r.bold = True
            p.add_run(texto)
        meta = doc.add_paragraph()
        meta.add_run(f"Entra: {et.entradas}\nSale: {et.salidas}\n"
                     f"Dónde vive: {' · '.join(et.modulos)}")
        for run in meta.runs:
            run.font.size = Pt(9)

    doc.add_page_break()
    pie = doc.add_paragraph(_pie(datos_demo))
    for run in pie.runs:
        run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PDF --------------------------------------------------------------------
def _latin1(texto: str) -> str:
    """FPDF con las fuentes del núcleo solo escribe latin-1.

    Un carácter fuera de ese rango —una comilla tipográfica pegada desde Word,
    un guion largo— hace explotar la generación entera con
    `UnicodeEncodeError`. Se reemplazan los sospechosos habituales por su
    equivalente y el resto se degrada, en vez de perder el PDF completo por un
    carácter.
    """
    reemplazos = {"—": "-", "–": "-", "“": '"', "”": '"', "‘": "'",
                  "’": "'", "…": "...", "·": "-", "→": "->", "≥": ">=",
                  "≤": "<=", "×": "x"}
    for viejo, nuevo in reemplazos.items():
        texto = texto.replace(viejo, nuevo)
    return texto.encode("latin-1", "replace").decode("latin-1")


def como_pdf(datos_demo: bool = True) -> bytes:
    """Mismo contenido, para imprimir y llevar."""
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    ancho = pdf.w - pdf.l_margin - pdf.r_margin

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(ancho, 8, _latin1(TITULO))
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(90, 105, 120)
    pdf.multi_cell(ancho, 5, _latin1(BAJADA))
    pdf.set_text_color(20, 30, 40)
    pdf.ln(3)

    for et in mt.etapas():
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(ancho, 6, _latin1(f"{et.orden}. {et.titulo}"))
        for clave in ("tecnico", "criollo", "por_que", "repercusion", "limites"):
            texto = getattr(et, clave)
            if not texto:
                continue
            pdf.set_font("Helvetica", "B", 9)
            pdf.multi_cell(ancho, 4.6, _latin1(_ETIQUETA[clave]))
            pdf.set_font("Helvetica", "", 9.5)
            pdf.multi_cell(ancho, 4.6, _latin1(texto))
            pdf.ln(0.8)
        pdf.set_font("Helvetica", "I", 8)
        pdf.set_text_color(105, 120, 135)
        pdf.multi_cell(ancho, 4, _latin1(
            f"Entra: {et.entradas} | Sale: {et.salidas} | "
            f"Donde vive: {' · '.join(et.modulos)}"))
        pdf.set_text_color(20, 30, 40)
        pdf.ln(3)

    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(105, 120, 135)
    pdf.multi_cell(ancho, 4, _latin1(_pie(datos_demo)))

    salida = pdf.output()
    return bytes(salida)
